# -*- coding: utf-8 -*-
"""Convert Markdown (OpenWorker-friendly) to a rough .docx for Word workflows.

Supports: ATX headings (#..######), paragraphs, unordered/ordered lists,
fenced code blocks, simple GFM tables, **bold** / *italic* / `code`.
Does NOT aim for perfect fidelity (no HTML, footnotes, nested lists, images).

Usage:
  python word-com/scripts/md_to_docx.py article.md --out article.docx
  python word-com/scripts/md_to_docx.py article.md --out article.docx --print
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor
except ImportError as exc:  # pragma: no cover
    raise SystemExit(
        "python-docx is required. Install with: pip install python-docx\n"
        f"ImportError: {exc}"
    ) from exc

INLINE_RE = re.compile(
    r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|\[[^\]]+\]\([^)]+\))"
)
HEADING_RE = re.compile(r"^(#{1,6})\s+(.+?)\s*$")
UL_RE = re.compile(r"^[-*+]\s+(.+)$")
OL_RE = re.compile(r"^(\d+)[.)]\s+(.+)$")
FENCE_RE = re.compile(r"^```")
TABLE_SEP_RE = re.compile(r"^\|?\s*:?-+:?\s*(\|\s*:?-+:?\s*)+\|?\s*$")


def _set_run_fonts(run, *, far_east: str = "宋体", ascii_font: str = "Times New Roman", size: float = 12) -> None:
    run.font.size = Pt(size)
    run.font.name = ascii_font
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), ascii_font)
    rFonts.set(qn("w:hAnsi"), ascii_font)
    rFonts.set(qn("w:eastAsia"), far_east)


def _add_inline_runs(paragraph, text: str, *, size: float = 12) -> None:
    if not text:
        return
    pos = 0
    for m in INLINE_RE.finditer(text):
        if m.start() > pos:
            run = paragraph.add_run(text[pos : m.start()])
            _set_run_fonts(run, size=size)
        token = m.group(0)
        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
            _set_run_fonts(run, size=size)
        elif token.startswith("*") and token.endswith("*"):
            run = paragraph.add_run(token[1:-1])
            run.italic = True
            _set_run_fonts(run, size=size)
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            _set_run_fonts(run, far_east="宋体", ascii_font="Consolas", size=size)
        elif token.startswith("["):
            # [text](url) → keep visible text only
            label = re.match(r"\[([^\]]+)\]", token)
            run = paragraph.add_run(label.group(1) if label else token)
            _set_run_fonts(run, size=size)
        else:
            run = paragraph.add_run(token)
            _set_run_fonts(run, size=size)
        pos = m.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        _set_run_fonts(run, size=size)


def _split_table_row(line: str) -> list[str]:
    s = line.strip().strip("|")
    return [c.strip() for c in s.split("|")]


def _is_table_row(line: str) -> bool:
    t = line.strip()
    return t.startswith("|") and t.count("|") >= 2


def md_to_document(md_text: str) -> Document:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    lines = md_text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    i = 0
    para_buf: list[str] = []

    def flush_para() -> None:
        nonlocal para_buf
        if not para_buf:
            return
        text = " ".join(x.strip() for x in para_buf if x.strip())
        para_buf = []
        if not text:
            return
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Pt(24)
        p.paragraph_format.line_spacing = 1.5
        _add_inline_runs(p, text, size=12)

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if FENCE_RE.match(stripped):
            flush_para()
            i += 1
            code_lines: list[str] = []
            while i < len(lines) and not FENCE_RE.match(lines[i].strip()):
                code_lines.append(lines[i])
                i += 1
            if i < len(lines):
                i += 1
            p = doc.add_paragraph()
            run = p.add_run("\n".join(code_lines))
            _set_run_fonts(run, far_east="宋体", ascii_font="Consolas", size=10)
            run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
            continue

        if _is_table_row(stripped) and i + 1 < len(lines) and TABLE_SEP_RE.match(lines[i + 1].strip()):
            flush_para()
            headers = _split_table_row(stripped)
            i += 2
            rows: list[list[str]] = []
            while i < len(lines) and _is_table_row(lines[i].strip()):
                rows.append(_split_table_row(lines[i].strip()))
                i += 1
            table = doc.add_table(rows=1 + len(rows), cols=len(headers))
            table.style = "Table Grid"
            for c, h in enumerate(headers):
                cell = table.rows[0].cells[c]
                cell.text = ""
                p = cell.paragraphs[0]
                _add_inline_runs(p, h, size=11)
            for r_i, row in enumerate(rows):
                for c, val in enumerate(row):
                    if c >= len(headers):
                        break
                    cell = table.rows[r_i + 1].cells[c]
                    cell.text = ""
                    p = cell.paragraphs[0]
                    _add_inline_runs(p, val, size=11)
            doc.add_paragraph()
            continue

        hm = HEADING_RE.match(stripped)
        if hm:
            flush_para()
            level = len(hm.group(1))
            text = hm.group(2).strip()
            # Strip inline markers for heading plain text (styles come from Word Heading N)
            plain = re.sub(r"[*_`]", "", text)
            p = doc.add_heading(plain, level=min(level, 9))
            if level == 1:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue

        um = UL_RE.match(stripped)
        if um:
            flush_para()
            p = doc.add_paragraph(style="List Bullet")
            _add_inline_runs(p, um.group(1), size=12)
            i += 1
            continue

        om = OL_RE.match(stripped)
        if om:
            flush_para()
            p = doc.add_paragraph(style="List Number")
            _add_inline_runs(p, om.group(2), size=12)
            i += 1
            continue

        if not stripped:
            flush_para()
            i += 1
            continue

        para_buf.append(stripped)
        i += 1

    flush_para()
    return doc


def convert_file(src: Path, out: Path) -> dict:
    text = src.read_text(encoding="utf-8")
    doc = md_to_document(text)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    return {
        "source": str(src.resolve()),
        "out": str(out.resolve()),
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "note": "Rough MD→DOCX for OpenWorker; school formatting still needs *-thesis-format.",
    }


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Markdown → rough .docx (OpenWorker bridge)")
    parser.add_argument("markdown", help="Input .md path")
    parser.add_argument("--out", required=True, help="Output .docx path")
    parser.add_argument("--print", action="store_true", dest="do_print", help="Print summary JSON")
    args = parser.parse_args(argv)

    src = Path(args.markdown)
    if not src.exists():
        print(f"ERROR: missing {src}", file=sys.stderr)
        return 1
    result = convert_file(src, Path(args.out))
    if args.do_print:
        import json

        print(json.dumps(result, ensure_ascii=False, indent=2))
    else:
        print("wrote", result["out"])
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
